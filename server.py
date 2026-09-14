import os
import sys
import io
import json
import time
from pathlib import Path
from typing import Any, List, Optional

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, File, UploadFile, HTTPException, Body, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, Field
from groq import Groq
from pypdf import PdfReader
from docx import Document

# ---------------------------------------------------------------------------
# API Key & Groq Model Config
# ---------------------------------------------------------------------------
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")
client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None
MODEL_NAME = "openai/gpt-oss-120b"

app = FastAPI(
    title="TalentMatch AI API",
    description="Automated Resume Parsing and Job-Resume Scoring API using Groq & Pydantic",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Pydantic Schemas
# ---------------------------------------------------------------------------
class JobD(BaseModel):
    role: str
    required_skills: List[str] = []
    preferred_skills: List[str] = []
    minimum_experience: Optional[float] = None
    education_requirements: List[str] = []
    responsibilities: List[str] = []

class Experience(BaseModel):
    company: Optional[str] = None
    role: Optional[str] = None
    duration: Optional[str] = None
    description: Optional[str] = None
    skills_used: List[str] = []

class Resume(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    total_experience_years: Optional[float] = None
    skills: List[Any] = []
    experiences: List[Experience] = []
    education: List[Any] = []
    projects: List[Any] = []
    certifications: List[Any] = []

class MatchResult(BaseModel):
    candidate_name: Optional[str] = None
    score: float
    details: dict = {}

class EvaluateTextRequest(BaseModel):
    candidate_name: Optional[str] = None
    resume_text: str

# ---------------------------------------------------------------------------
# Default Amazon SDE-I Job Description
# ---------------------------------------------------------------------------
DEFAULT_JOB_DESCRIPTION = """
Software Development Engineer I (SDE-I) at Amazon.
Key job responsibilities:
- Collaborate and communicate effectively with cross-disciplinary Amazonians to design, build, and operate innovative products.
- Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
- Participate in code reviews and contribute to technical documentation.
- Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
- Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity.

Basic Qualifications:
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or in the process of obtaining a bachelor's degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age or older

Preferred Qualifications:
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with AI tools for development productivity
- Cloud platforms (preferably AWS)
- Database systems (SQL and NoSQL)
- Contributing to open-source projects
- Version control systems (Git/GitHub)
- Debugging and troubleshooting complex systems
"""

active_job_description = DEFAULT_JOB_DESCRIPTION
active_job_obj: Optional[JobD] = None

# Pre-seeded Day 5 Candidates
DEFAULT_CANDIDATES = [
    {
        "id": "ashish-raj",
        "name": "Ashish Raj",
        "score": 70.0,
        "rank": 1,
        "role": "SDE Aspirant · Cloud & DevOps Engineer",
        "email": "ashish.raj@example.com",
        "phone": "+91 98765 43210",
        "location": "Delhi NCR, India",
        "fileSource": "Ashish Raj 24PCS007 (1) - Ashish Raj.pdf",
        "matchingSkills": [
            "C++", "Python", "Data Structures", "Algorithms", "AWS", "GitHub", 
            "CI/CD", "Docker", "Kubernetes", "Jenkins", "Ansible", "Communication", "Problem Solving"
        ],
        "missingSkills": [
            "Object-oriented design principles", "Database systems (SQL/NoSQL)", 
            "Open-source contributions", "Explicit SDLC understanding"
        ],
        "experienceRequirementMet": True,
        "overallMatchPercentage": 70,
        "finalVerdict": "Good fit for an entry‑level software role with standout DevOps automation and programming foundations, but should strengthen OOP design patterns and database experience to fully meet job expectations.",
        "education": ["Master of Computer Applications (MCA) — 2024-2026", "Bachelor of Computer Applications (BCA) — 2021-2024 (CGPA: 8.6/10)"],
        "experiences": [
            {
                "role": "Cloud & DevOps Intern",
                "company": "TechSolutions Cloud Labs",
                "duration": "6 Months",
                "description": "Automated microservices deployment pipelines using Docker, Kubernetes, Jenkins, and AWS EC2/S3. Reduced deployment lead time by 35%.",
                "skillsUsed": ["Docker", "Kubernetes", "AWS", "CI/CD", "Jenkins", "Ansible"]
            }
        ],
        "projects": [
            {
                "title": "Cloud Infrastructure Automation & CI/CD Pipeline",
                "tech": ["AWS", "Docker", "Kubernetes", "GitHub Actions"],
                "description": "Designed multi-tier resilient microservices architecture on AWS with automated canary rollouts and Prometheus monitoring."
            },
            {
                "title": "Smart Vision IoT & Microcontroller Pipeline",
                "tech": ["Python", "OpenCV", "C++", "Arduino"],
                "description": "Real-time edge computing vision pipeline for anomaly detection with low-latency data streaming."
            }
        ],
        "status": "Top Match",
        "badgeTier": "tier-gold"
    },
    {
        "id": "abhay-singh",
        "name": "Abhay Pratap Singh",
        "score": 62.0,
        "rank": 2,
        "role": "Software Development Engineer Trainee",
        "email": "abhay.singh@example.com",
        "phone": "+91 91234 56789",
        "location": "Bengaluru, India",
        "fileSource": "abhay resume new - Abhay Singh.pdf",
        "matchingSkills": ["C++", "Data Structures & Algorithms", "SQL", "Git/GitHub"],
        "missingSkills": [
            "Object-oriented design principles", "AI tool experience", "Cloud platforms (AWS)", 
            "Open‑source contributions", "Advanced version‑control practices", "Debugging & troubleshooting", 
            "SDLC understanding", "Strong written & verbal communication"
        ],
        "experienceRequirementMet": True,
        "overallMatchPercentage": 62,
        "finalVerdict": "Meets basic education and core required programming skills. Demonstrates solid fundamentals in C++ and relational databases, with room for growth in cloud tools and system design.",
        "education": ["B.Tech in Computer Science & Engineering — 2020-2024 (CGPA: 8.2/10)"],
        "experiences": [
            {
                "role": "Software Engineering Intern",
                "company": "Apex Infotech",
                "duration": "3 Months",
                "description": "Optimized relational database queries, resolved bottlenecks, and implemented core algorithmic routines in C++.",
                "skillsUsed": ["C++", "SQL", "Data Structures", "Git"]
            }
        ],
        "projects": [
            {
                "title": "High Performance DSA Algorithmic Engine",
                "tech": ["C++", "STL", "Memory Management"],
                "description": "High-throughput engine implementing graph traversal and dynamic programming optimizations for resource allocation."
            }
        ],
        "status": "Recommended",
        "badgeTier": "tier-silver"
    },
    {
        "id": "anshit-verma",
        "name": "Anshit Verma",
        "score": 55.0,
        "rank": 3,
        "role": "Junior Software Developer",
        "email": "anshit.verma@example.com",
        "phone": "+91 97654 32109",
        "location": "Noida, India",
        "fileSource": "anshit verma resume word - Anshit Verma.docx",
        "matchingSkills": ["C++", "Python (basic)", "MySQL (basic)", "Problem-solving", "Adaptability", "Teamwork", "Communication"],
        "missingSkills": [
            "Data structures implementation", "Algorithm development", "Object-oriented design principles", 
            "Technical internship experience", "AI development tools", "AWS or other cloud platforms", 
            "Version control systems (Git)", "Debugging complex systems", "Explicit SDLC knowledge"
        ],
        "experienceRequirementMet": True,
        "overallMatchPercentage": 55,
        "finalVerdict": "The candidate possesses basic programming and analytical abilities but lacks several core technical and preferred qualifications needed for the SDE-I role.",
        "education": ["B.Tech in Computer Science & Engineering — 2021-2025"],
        "experiences": [
            {
                "role": "Academic Project Lead",
                "company": "University Tech Club",
                "duration": "Academic",
                "description": "Collaborated in student teams building modular utility applications and database-backed management scripts.",
                "skillsUsed": ["Python", "MySQL", "Teamwork"]
            }
        ],
        "projects": [
            {
                "title": "Store Inventory Management System",
                "tech": ["Python", "MySQL"],
                "description": "Desktop application for tracking store supplies, generating stock depletion alerts, and monthly sales logs."
            }
        ],
        "status": "Under Review",
        "badgeTier": "tier-bronze"
    },
    {
        "id": "priyanshu-singh",
        "name": "Priyanshu Singh",
        "score": 45.0,
        "rank": 4,
        "role": "Full Stack Web Developer",
        "email": "priyanshu.singh@example.com",
        "phone": "+91 99887 76655",
        "location": "Lucknow, India",
        "fileSource": "Resume for Freshers - Priyanshu Singh.docx",
        "matchingSkills": ["Java", "SQL", "MySQL", "MongoDB", "Git", "GitHub", "React.js", "Node.js", "Express.js", "Redux Toolkit"],
        "missingSkills": [
            "Data structures implementation", "Algorithm development", "Object-oriented design principles", 
            "Experience with AI tools", "AWS/cloud platforms", "Open-source contributions", 
            "Debugging & troubleshooting", "SDLC knowledge"
        ],
        "experienceRequirementMet": True,
        "overallMatchPercentage": 45,
        "finalVerdict": "Candidate shows solid full‑stack JavaScript/React development capabilities and relevant degree, but lacks the core distributed systems, cloud (AWS), and advanced DSA expected for Amazon SDE-I.",
        "education": ["B.Tech in Computer Science & Engineering — 2021-2025"],
        "experiences": [
            {
                "role": "Full Stack Web Developer Intern",
                "company": "Freelance / Digital Studio",
                "duration": "4 Months",
                "description": "Constructed client web portals with React.js, Express, and MongoDB. Handled state management with Redux.",
                "skillsUsed": ["React.js", "Node.js", "Express.js", "MongoDB", "Redux"]
            }
        ],
        "projects": [
            {
                "title": "Modern E-Commerce Storefront",
                "tech": ["React.js", "Node.js", "MongoDB", "TailwindCSS"],
                "description": "Full featured commerce web application with product filters, shopping cart, and JWT-authenticated checkout flow."
            }
        ],
        "status": "Needs Skill Upskilling",
        "badgeTier": "tier-muted"
    }
]

candidates_db = list(DEFAULT_CANDIDATES)

# ---------------------------------------------------------------------------
# LLM Helper Functions with Auto-Retry
# ---------------------------------------------------------------------------
def call_groq_with_retry(messages: list, response_format: dict, max_retries: int = 5):
    if not client:
        raise ValueError("GROQ_API_KEY is not configured on the server.")

    for attempt in range(max_retries):
        try:
            return client.chat.completions.create(
                model=MODEL_NAME,
                messages=messages,
                response_format=response_format
            )
        except Exception as e:
            err_str = str(e).lower()
            if "rate_limit" in err_str or "429" in err_str:
                if attempt == max_retries - 1:
                    raise
                wait_sec = 6 * (attempt + 1)
                print(f"[Groq RateLimit] Sleeping {wait_sec}s before retry (attempt {attempt + 1}/{max_retries})...")
                time.sleep(wait_sec)
            else:
                raise

def get_parsed_job_object() -> JobD:
    global active_job_obj
    if active_job_obj:
        return active_job_obj

    if not client:
        return JobD(
            role="Software Development Engineer I",
            required_skills=["C++", "Java", "Python", "Data Structures", "Algorithms", "OOP"],
            preferred_skills=["AWS", "Docker", "Kubernetes", "Git", "SQL"],
            minimum_experience=0.0,
            education_requirements=["Bachelor's degree in CS, CE, or STEM"],
            responsibilities=["Design distributed microservices", "Collaborate with cross-functional teams"]
        )

    job_schema = JobD.model_json_schema()
    system_prompt = f"""You are an expert HR assistant.
Extract structured information from the job description.
Return ONLY valid JSON matching this schema:
{job_schema}
If minimum experience is not mentioned, return 0.0.
"""
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Analyze this job description:\n{active_job_description}"}
    ]
    resp = call_groq_with_retry(messages, {"type": "json_object"})
    data = json.loads(resp.choices[0].message.content)
    active_job_obj = JobD(**data)
    return active_job_obj

def parse_resume_text(resume_text: str) -> Resume:
    resume_schema = Resume.model_json_schema()
    system_prompt = f"""You are an expert resume parser.
Extract information from the resume based on its semantic meaning.
Return ONLY valid JSON matching this schema:
{resume_schema}
Rules:
1. Do not invent information.
2. If value not available, return null.
3. If a list has no information, return empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across entire resume.
"""
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Parse the following resume:\n{resume_text}"}
    ]
    resp = call_groq_with_retry(messages, {"type": "json_object"})
    data = json.loads(resp.choices[0].message.content)
    return Resume(**data)

def score_resume_against_job(job: JobD, resume: Resume) -> dict:
    prompt = f"""You are an HR recruiter.
Compare the candidate's resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Return JSON with this exact structure:
{{
  "candidate_name": "string",
  "score": float (0 to 100),
  "matching_skills": ["list", "of", "matching", "skills"],
  "missing_important_skills": ["list", "of", "missing", "skills"],
  "experience_requirement_met": true/false,
  "final_verdict": "string (concise recruiter assessment)"
}}
Keep response concise and easy to read.
"""
    messages = [{"role": "user", "content": prompt}]
    resp = call_groq_with_retry(messages, {"type": "json_object"})
    return json.loads(resp.choices[0].message.content)

# ---------------------------------------------------------------------------
# File Extraction Helpers
# ---------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in reader.pages:
        t = page.extract_text()
        if t:
            text += t + "\n"
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    text = ""
    for p in doc.paragraphs:
        if p.text.strip():
            text += p.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

# ---------------------------------------------------------------------------
# API Endpoints
# ---------------------------------------------------------------------------
@app.get("/api/health")
def get_health():
    return {
        "status": "healthy",
        "has_groq_key": bool(GROQ_API_KEY),
        "model": MODEL_NAME,
        "candidates_count": len(candidates_db)
    }

@app.get("/api/job-description")
def get_job_description():
    job_obj = get_parsed_job_object()
    return {
        "raw_text": active_job_description,
        "structured": job_obj.model_dump()
    }

@app.post("/api/job-description")
def update_job_description(body: dict = Body(...)):
    global active_job_description, active_job_obj
    new_text = body.get("job_description", "")
    if not new_text.strip():
        raise HTTPException(status_code=400, detail="Job description text cannot be empty.")
    active_job_description = new_text
    active_job_obj = None # will be re-parsed on next request
    return {"message": "Job description updated successfully"}

@app.get("/api/candidates")
def get_candidates():
    # Return candidates sorted by score
    return sorted(candidates_db, key=lambda c: c["score"], reverse=True)

@app.post("/api/reset-candidates")
def reset_candidates():
    global candidates_db
    candidates_db = list(DEFAULT_CANDIDATES)
    return {"message": "Candidates reset to default Day 5 cohort", "candidates": candidates_db}

@app.post("/api/upload-resume")
async def upload_resume(file: UploadFile = File(...)):
    filename = file.filename or "uploaded_resume.pdf"
    ext = Path(filename).suffix.lower()
    
    if ext not in [".pdf", ".docx"]:
        raise HTTPException(status_code=400, detail="Only .pdf and .docx files are supported.")
    
    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    try:
        if ext == ".pdf":
            resume_text = extract_text_from_pdf(file_bytes)
        else:
            resume_text = extract_text_from_docx(file_bytes)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Failed to extract text from document: {str(e)}")

    if not resume_text.strip():
        raise HTTPException(status_code=422, detail="No readable text could be extracted from this document.")

    # Parse and Score
    job_obj = get_parsed_job_object()
    try:
        parsed_resume = parse_resume_text(resume_text)
        time.sleep(2) # brief pause to respect Groq TPM
        score_data = score_resume_against_job(job_obj, parsed_resume)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI evaluation failed: {str(e)}")

    score_val = float(score_data.get("score", 50.0))
    cand_name = score_data.get("candidate_name") or parsed_resume.name or Path(filename).stem
    
    new_candidate = {
        "id": f"upload-{int(time.time()*1000)}",
        "name": cand_name,
        "score": score_val,
        "rank": 1, # will be updated after sort
        "role": "Candidate (Uploaded Resume)",
        "email": parsed_resume.email or "not_provided@candidate.com",
        "phone": parsed_resume.phone or "Not provided",
        "location": "Uploaded Profile",
        "fileSource": filename,
        "matchingSkills": score_data.get("matching_skills", []),
        "missingSkills": score_data.get("missing_important_skills", []),
        "experienceRequirementMet": bool(score_data.get("experience_requirement_met", True)),
        "overallMatchPercentage": int(score_val),
        "finalVerdict": score_data.get("final_verdict", "Candidate evaluated successfully."),
        "education": [str(e) for e in parsed_resume.education] if parsed_resume.education else ["Degree details not specified"],
        "experiences": [exp.model_dump() for exp in parsed_resume.experiences],
        "projects": parsed_resume.projects,
        "status": "Top Match" if score_val >= 70 else "Recommended" if score_val >= 60 else "Needs Review",
        "badgeTier": "tier-gold" if score_val >= 70 else "tier-silver" if score_val >= 60 else "tier-bronze" if score_val >= 50 else "tier-muted"
    }

    candidates_db.append(new_candidate)
    candidates_db.sort(key=lambda c: c["score"], reverse=True)
    for i, c in enumerate(candidates_db):
        c["rank"] = i + 1

    return {"candidate": new_candidate, "all_candidates": candidates_db}

@app.post("/api/analyze-single-resume")
async def analyze_single_resume(
    file: Optional[UploadFile] = File(None),
    resume_text: Optional[str] = Form(None),
    candidate_name: Optional[str] = Form(None),
    custom_job_description: Optional[str] = Form(None),
    save_to_cohort: Optional[bool] = Form(True)
):
    source_name = "Text Input"
    text_content = ""

    if file and file.filename:
        source_name = file.filename
        ext = Path(source_name).suffix.lower()
        if ext not in [".pdf", ".docx"]:
            raise HTTPException(status_code=400, detail="Only .pdf and .docx documents are supported.")
        file_bytes = await file.read()
        if not file_bytes:
            raise HTTPException(status_code=400, detail="Uploaded file is empty.")
        if ext == ".pdf":
            text_content = extract_text_from_pdf(file_bytes)
        else:
            text_content = extract_text_from_docx(file_bytes)
    elif resume_text and resume_text.strip():
        text_content = resume_text.strip()
    else:
        raise HTTPException(status_code=400, detail="Please upload a resume file (.pdf/.docx) or provide resume text.")

    if not text_content.strip():
        raise HTTPException(status_code=422, detail="No readable text found in the provided resume.")

    # Target Job Specification
    if custom_job_description and custom_job_description.strip():
        job_schema = JobD.model_json_schema()
        system_prompt = f"Extract structured job requirements. Return ONLY valid JSON matching this schema:\n{job_schema}\nIf min experience missing, return 0.0."
        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": f"Analyze this job description:\n{custom_job_description.strip()}"}
        ]
        resp = call_groq_with_retry(messages, {"type": "json_object"})
        job_obj = JobD(**json.loads(resp.choices[0].message.content))
        time.sleep(1)
    else:
        job_obj = get_parsed_job_object()

    # Parse and Score
    try:
        parsed_resume = parse_resume_text(text_content)
        time.sleep(2)
        score_data = score_resume_against_job(job_obj, parsed_resume)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI evaluation failed: {str(e)}")

    score_val = float(score_data.get("score", 50.0))
    final_cand_name = candidate_name or score_data.get("candidate_name") or parsed_resume.name or Path(source_name).stem

    candidate_obj = {
        "id": f"single-{int(time.time()*1000)}",
        "name": final_cand_name,
        "score": score_val,
        "rank": 1,
        "role": f"Candidate for {job_obj.role}",
        "email": parsed_resume.email or "candidate@email.com",
        "phone": parsed_resume.phone or "+91 00000 00000",
        "location": "Candidate Profile",
        "fileSource": source_name,
        "matchingSkills": score_data.get("matching_skills", []),
        "missingSkills": score_data.get("missing_important_skills", []),
        "experienceRequirementMet": bool(score_data.get("experience_requirement_met", True)),
        "overallMatchPercentage": int(score_val),
        "finalVerdict": score_data.get("final_verdict", "Candidate evaluated successfully."),
        "education": [str(e) for e in parsed_resume.education] if parsed_resume.education else ["Degree info extracted"],
        "experiences": [exp.model_dump() for exp in parsed_resume.experiences],
        "projects": parsed_resume.projects,
        "status": "Top Match" if score_val >= 70 else "Recommended" if score_val >= 60 else "Needs Review",
        "badgeTier": "tier-gold" if score_val >= 70 else "tier-silver" if score_val >= 60 else "tier-bronze"
    }

    if save_to_cohort:
        candidates_db.append(candidate_obj)
        candidates_db.sort(key=lambda c: c["score"], reverse=True)
        for i, c in enumerate(candidates_db):
            c["rank"] = i + 1

    return {
        "candidate": candidate_obj,
        "target_role": job_obj.role,
        "matching_skills": candidate_obj["matchingSkills"],
        "missing_skills": candidate_obj["missingSkills"],
        "score": score_val,
        "verdict": candidate_obj["finalVerdict"],
        "experience_met": candidate_obj["experienceRequirementMet"],
        "total_cohort_count": len(candidates_db)
    }


@app.post("/api/evaluate-text")
def evaluate_text(req: EvaluateTextRequest):
    if not req.resume_text.strip():
        raise HTTPException(status_code=400, detail="Resume text cannot be empty.")

    job_obj = get_parsed_job_object()
    try:
        parsed_resume = parse_resume_text(req.resume_text)
        time.sleep(2)
        score_data = score_resume_against_job(job_obj, parsed_resume)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI evaluation failed: {str(e)}")

    score_val = float(score_data.get("score", 50.0))
    cand_name = req.candidate_name or score_data.get("candidate_name") or parsed_resume.name or "Candidate"

    new_candidate = {
        "id": f"text-{int(time.time()*1000)}",
        "name": cand_name,
        "score": score_val,
        "rank": 1,
        "role": "Candidate (Direct Input)",
        "email": parsed_resume.email or "candidate@email.com",
        "phone": parsed_resume.phone or "+91 00000 00000",
        "location": "Direct Evaluation",
        "fileSource": f"{cand_name.replace(' ', '_')}_text.txt",
        "matchingSkills": score_data.get("matching_skills", []),
        "missingSkills": score_data.get("missing_important_skills", []),
        "experienceRequirementMet": bool(score_data.get("experience_requirement_met", True)),
        "overallMatchPercentage": int(score_val),
        "finalVerdict": score_data.get("final_verdict", "Candidate evaluated successfully."),
        "education": [str(e) for e in parsed_resume.education] if parsed_resume.education else ["Extracted from text"],
        "experiences": [exp.model_dump() for exp in parsed_resume.experiences],
        "projects": parsed_resume.projects,
        "status": "Top Match" if score_val >= 70 else "Recommended" if score_val >= 60 else "Needs Review",
        "badgeTier": "tier-gold" if score_val >= 70 else "tier-silver" if score_val >= 60 else "tier-bronze"
    }

    candidates_db.append(new_candidate)
    candidates_db.sort(key=lambda c: c["score"], reverse=True)
    for i, c in enumerate(candidates_db):
        c["rank"] = i + 1

    return {"candidate": new_candidate, "all_candidates": candidates_db}

# ---------------------------------------------------------------------------
# Static Frontend Serving for Production
# ---------------------------------------------------------------------------
frontend_dist = Path(__file__).parent / "frontend" / "dist"
if frontend_dist.exists():
    app.mount("/assets", StaticFiles(directory=str(frontend_dist / "assets")), name="assets")

    @app.get("/{full_path:path}")
    async def serve_spa(full_path: str):
        # Serve exact file if exists, else index.html for SPA routing
        target = frontend_dist / full_path
        if target.exists() and target.is_file():
            return FileResponse(target)
        return FileResponse(frontend_dist / "index.html")

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    print(f"Starting TalentMatch AI Server on http://0.0.0.0:{port}")
    uvicorn.run("server:app", host="0.0.0.0", port=port, reload=False)
